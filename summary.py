"""
The part that does the comparing.

After the pages are filed, the app reads the numbers off each property's
section (ledger.py), keeps one small JSON file per month under
    <destination>/Monthly Summary/data/<YYYY-MM>.json
and rebuilds
    <destination>/Monthly Summary/Monthly Summary.xlsx
from every month on record. The workbook has four sheets:

  What changed   the point of it all — one line per thing worth a look,
                 this month against the previous month on record, plus a
                 few checks that don't need a prior month (negative cash,
                 tenant past due, vacancy, reversal entries that don't net
                 to zero, a management fee out of line with the others)
  This month     one row per property with the key figures
  History        property × month grids for the figures worth trending
  Ledger         every transaction, every month, filterable

Nothing here is a judgement about whether a charge is *right* — it only
points at what is different or unusual, so a person can look there first.
"""
from __future__ import annotations

import json
import re
import statistics
from dataclasses import dataclass, field
from datetime import datetime
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill, Border, Side
from openpyxl.utils import get_column_letter

from ledger import read_property_month

SUMMARY_DIR = "Monthly Summary"
DATA_DIR = "data"
WORKBOOK = "Monthly Summary.xlsx"

MONTH_WORDS = ("january february march april may june july august september october "
               "november december jan feb mar apr jun jul aug sep sept oct nov dec").split()


# ------------------------------------------------------------------ data model

@dataclass
class Change:
    property: str          # folder name, e.g. "004-01 123 Main Street"
    severity: str          # "attention" | "note"
    what: str              # short label
    detail: str            # plain-English sentence
    this: float | str | None = None
    last: float | str | None = None


@dataclass
class SummaryResult:
    workbook: Path
    month: str                   # the month the comparison is FOR (the latest on record)
    month_label: str
    compared_to: str | None      # label of the prior month used, or None
    changes: list[Change] = field(default_factory=list)
    months_on_record: int = 0
    ran_month_label: str = ""    # the month that was just processed
    ran_is_latest: bool = True   # False when an earlier month was run to fill in history

    @property
    def attention(self) -> list[Change]:
        return [c for c in self.changes if c.severity == "attention"]

    @property
    def notes(self) -> list[Change]:
        return [c for c in self.changes if c.severity == "note"]


# ------------------------------------------------------------------ collecting

def collect_month(plan, reader) -> dict:
    props = {}
    for s in plan.sections:
        if s.code is None:
            continue
        d = read_property_month(reader, s.first_page, s.last_page)
        d.update({"code": s.code, "address": s.address, "duplicate": s.duplicate_of,
                  "first_page": s.first_page, "last_page": s.last_page})
        props[s.folder_name] = d
    return {"month": plan.month, "month_label": plan.month_label,
            "source": plan.source.name, "generated": datetime.now().isoformat(timespec="seconds"),
            "properties": props}


def _data_dir(destination: Path) -> Path:
    return Path(destination) / SUMMARY_DIR / DATA_DIR


def save_month(destination: Path, month_data: dict) -> Path:
    d = _data_dir(destination)
    d.mkdir(parents=True, exist_ok=True)
    p = d / f"{month_data['month']}.json"
    p.write_text(json.dumps(month_data, indent=1), encoding="utf-8")
    return p


def load_months(destination: Path) -> dict[str, dict]:
    out = {}
    d = _data_dir(destination)
    if d.is_dir():
        for p in sorted(d.glob("*.json")):
            try:
                m = json.loads(p.read_text(encoding="utf-8"))
                out[m["month"]] = m
            except (OSError, ValueError, KeyError):
                continue
    return out


# ------------------------------------------------------------------ derived figures

def _norm(text: str) -> str:
    """Strip the bits that legitimately change month to month (dates, account
    numbers, month names) so the same recurring charge gets the same key."""
    t = (text or "").lower()
    t = re.sub(r"\d[\d/,.\-]*", " ", t)
    t = " ".join(w for w in t.split() if w not in MONTH_WORDS)
    return re.sub(r"\s+", " ", t).strip(" -")


def figures(p: dict) -> dict:
    led = p.get("ledger") or []
    def s(pred, key):
        return round(sum((l.get(key) or 0) for l in led if pred(l)), 2)
    desc = lambda l: (l.get("description") or "").lower()
    typ = lambda l: (l.get("type") or "").lower()
    # "Rent Income - August 2026", plain "Rent Income", and Section 8 rent all count as rent.
    rent_received = s(lambda l: desc(l).startswith("rent income") or "section 8 rent" in desc(l), "income")
    mgmt_fee = s(lambda l: desc(l).startswith("management fee"), "expense")
    owner_payment = s(lambda l: desc(l).startswith("owner distribution"), "expense")
    reversals = [l for l in led if "revers" in typ(l)]
    rev_net = round(sum((l.get("income") or 0) - (l.get("expense") or 0) for l in reversals), 2)
    units = p.get("units") or []
    # A property with several units (a duplex) has one rent-roll row per unit:
    # add the money up, list every tenant, and call it "Current" only if all are.
    tenants = [u.get("tenant") for u in units if u.get("tenant")]
    statuses = [u.get("status") for u in units if u.get("status")]
    money_rows = [u for u in units if u.get("rent") is not None or u.get("recurring") is not None]
    sched = round(sum((u.get("rent") or 0) + (u.get("recurring") or 0) for u in money_rows), 2) if money_rows else None
    past_rows = [u.get("past_due") for u in units if u.get("past_due") is not None]
    u = {
        "tenant": ", ".join(tenants),
        "status": ("Current" if all(s.lower() == "current" for s in statuses) else
                   ", ".join(sorted(set(s for s in statuses if s.lower() != "current")))) if statuses else "",
        "rent": round(sum(u.get("rent") or 0 for u in money_rows), 2) if money_rows else None,
        "recurring": round(sum(u.get("recurring") or 0 for u in money_rows), 2) if money_rows else None,
        "past_due": round(sum(past_rows), 2) if past_rows else None,
    }
    # The manager's fee is a percentage of the *scheduled* rent (rent roll), whether or
    # not the tenant paid it all; fall back to rent received when there's no rent roll.
    fee_base = sched if sched else rent_received
    return {
        "tenant": u.get("tenant") or "",
        "status": u.get("status") or "",
        "scheduled_rent": sched,
        "rent_roll_rent": u.get("rent"),
        "recurring": u.get("recurring"),
        "past_due": u.get("past_due"),
        "rent_received": rent_received,
        "total_income": p.get("total_income"),
        "other_income": round((p.get("total_income") or 0) - rent_received, 2),
        "mgmt_fee": mgmt_fee,
        "mgmt_pct": round(100 * mgmt_fee / fee_base, 1) if fee_base else None,
        "owner_payment": owner_payment,
        "total_expense": p.get("total_expense"),
        "other_expense": round((p.get("total_expense") or 0) - mgmt_fee - owner_payment, 2),
        "net_income": p.get("is_net_income"),
        "net_income_ytd": p.get("is_net_income_ytd"),
        "beginning_cash": p.get("beginning_cash"),
        "ending_cash": p.get("ending_cash"),
        "bills_due": p.get("bills_due_total"),
        "reversal_count": len(reversals),
        "reversal_net": rev_net,
        "pages": p.get("pages"),
    }


def expense_keys(p: dict) -> dict[str, dict]:
    """Recurring-style expense lines keyed by normalised payee+description."""
    out: dict[str, dict] = {}
    for l in p.get("ledger") or []:
        if not l.get("expense"):
            continue
        d = (l.get("description") or "").lower()
        if d.startswith(("management fee", "owner distribution")) or "revers" in (l.get("type") or "").lower():
            continue
        key = _norm(l.get("payee", "")) + " | " + _norm(l.get("description", ""))
        if key in out:
            out[key]["amount"] = round(out[key]["amount"] + l["expense"], 2)
            out[key]["count"] += 1
        else:
            out[key] = {"payee": l.get("payee", ""), "description": l.get("description", ""),
                        "amount": l["expense"], "count": 1}
    return out


def _fmt(v) -> str:
    if v is None:
        return "—"
    if isinstance(v, (int, float)):
        return f"${v:,.2f}" if v >= 0 else f"-${-v:,.2f}"
    return str(v)


# ------------------------------------------------------------------ the checks

def compare(cur: dict, prev: dict | None) -> list[Change]:
    changes: list[Change] = []
    cp = cur["properties"]
    pp = prev["properties"] if prev else {}
    figs = {k: figures(v) for k, v in cp.items()}
    pfigs = {k: figures(v) for k, v in pp.items()}
    add = lambda *a, **k: changes.append(Change(*a, **k))

    # ---- checks that need no prior month
    pcts = [f["mgmt_pct"] for f in figs.values() if f["mgmt_pct"] is not None and f["mgmt_fee"] > 0]
    typical_pct = statistics.median(pcts) if len(pcts) >= 3 else None
    for name, f in figs.items():
        if f["ending_cash"] is not None and f["ending_cash"] < 0:
            add(name, "attention", "Negative cash balance",
                f"Ending cash balance is {_fmt(f['ending_cash'])}.", f["ending_cash"])
        if f["past_due"] is not None and f["past_due"] > 5.0:
            last = pfigs.get(name, {}).get("past_due")
            if last is None or last <= 5.0:
                add(name, "attention", "Tenant past due",
                    f"{f['tenant'] or 'Tenant'} is {_fmt(f['past_due'])} past due.", f["past_due"], last)
            elif abs(f["past_due"] - last) > 0.5:
                word = "up from" if f["past_due"] > last else "down from"
                add(name, "attention", "Past due changed",
                    f"{f['tenant'] or 'Tenant'} is {_fmt(f['past_due'])} past due, {word} {_fmt(last)}.", f["past_due"], last)
        if f["status"] and f["status"].lower() != "current":
            add(name, "attention", "Unit not current", f"Rent roll status is \"{f['status']}\".", f["status"])
        if f["reversal_count"] and abs(f["reversal_net"]) > 0.005:
            add(name, "attention", "Reversals don't net to zero",
                f"{f['reversal_count']} reversal entries net to {_fmt(f['reversal_net'])} instead of $0.00.",
                f["reversal_net"])
        if typical_pct is not None and f["mgmt_pct"] is not None and abs(f["mgmt_pct"] - typical_pct) > 1.0:
            add(name, "attention", "Management fee out of line",
                f"Management fee is {f['mgmt_pct']:.1f}% of scheduled rent; the typical property pays {typical_pct:.1f}%.",
                f"{f['mgmt_pct']:.1f}%", f"{typical_pct:.1f}%")
        if f["rent_received"] == 0 and f["status"].lower() == "current" and (f["scheduled_rent"] or 0) > 0:
            add(name, "attention", "No rent received",
                f"No rent income posted this month; rent roll expects {_fmt(f['scheduled_rent'])}.", 0, f["scheduled_rent"])

    if not prev:
        return changes

    # ---- month-over-month
    for name in sorted(set(pp) - set(cp)):
        add(name, "attention", "Property missing", f"Was in {prev['month_label']} but not in {cur['month_label']}.")
    for name in sorted(set(cp) - set(pp)):
        add(name, "note", "New property", f"Not in {prev['month_label']}; first seen in {cur['month_label']}.")

    for name in sorted(set(cp) & set(pp)):
        f, g = figs[name], pfigs[name]
        if f["tenant"] and g["tenant"] and _norm(f["tenant"]) != _norm(g["tenant"]):
            add(name, "attention", "Tenant changed", f"Tenant is now {f['tenant']} (was {g['tenant']}).", f["tenant"], g["tenant"])
        if f["scheduled_rent"] is not None and g["scheduled_rent"] is not None and abs(f["scheduled_rent"] - g["scheduled_rent"]) > 0.5:
            add(name, "attention", "Rent roll changed",
                f"Scheduled rent plus recurring charges is {_fmt(f['scheduled_rent'])} (was {_fmt(g['scheduled_rent'])}).",
                f["scheduled_rent"], g["scheduled_rent"])
        if abs(f["rent_received"] - g["rent_received"]) > 0.5:
            sev = "attention" if f["rent_received"] < g["rent_received"] else "note"
            add(name, sev, "Rent received changed",
                f"Rent income {_fmt(f['rent_received'])} (was {_fmt(g['rent_received'])}).", f["rent_received"], g["rent_received"])
        if abs(f["mgmt_fee"] - g["mgmt_fee"]) > 0.5 and not (f["mgmt_pct"] and g["mgmt_pct"] and abs(f["mgmt_pct"] - g["mgmt_pct"]) <= 0.2):
            add(name, "note", "Management fee changed",
                f"Management fee {_fmt(f['mgmt_fee'])} (was {_fmt(g['mgmt_fee'])}).", f["mgmt_fee"], g["mgmt_fee"])
        if g["owner_payment"] and abs(f["owner_payment"] - g["owner_payment"]) > max(50.0, 0.25 * g["owner_payment"]):
            add(name, "note", "Owner payment changed a lot",
                f"Owner payment {_fmt(f['owner_payment'])} (was {_fmt(g['owner_payment'])}).", f["owner_payment"], g["owner_payment"])
        if f["bills_due"] is not None and g["bills_due"] is not None and abs(f["bills_due"] - g["bills_due"]) > 1.0:
            add(name, "note", "Bills due changed", f"Bills due {_fmt(f['bills_due'])} (was {_fmt(g['bills_due'])}).", f["bills_due"], g["bills_due"])

        ce, pe = expense_keys(cp[name]), expense_keys(pp[name])
        for key in sorted(set(ce) & set(pe)):
            a, b = ce[key]["amount"], pe[key]["amount"]
            if abs(a - b) > 0.5:
                add(name, "note", "Charge amount changed",
                    f"{ce[key]['payee']}: {ce[key]['description']} — {_fmt(a)} (was {_fmt(b)}).", a, b)
        for key in sorted(set(ce) - set(pe)):
            add(name, "note", "New charge",
                f"{ce[key]['payee']}: {ce[key]['description']} — {_fmt(ce[key]['amount'])}; nothing like it last month.",
                ce[key]["amount"])
        for key in sorted(set(pe) - set(ce)):
            add(name, "note", "Charge not seen this month",
                f"{pe[key]['payee']}: {pe[key]['description']} — was {_fmt(pe[key]['amount'])} last month, nothing this month.",
                None, pe[key]["amount"])
    order = {"attention": 0, "note": 1}
    changes.sort(key=lambda c: (order[c.severity], c.property, c.what))
    return changes


# ------------------------------------------------------------------ the workbook

HEAD_FILL = PatternFill("solid", fgColor="0F766E")
HEAD_FONT = Font(bold=True, color="FFFFFF")
ATTN_FILL = PatternFill("solid", fgColor="FDE2E1")
NOTE_FILL = PatternFill("solid", fgColor="FFF4D6")
TITLE_FONT = Font(bold=True, size=14)
SUB_FONT = Font(italic=True, color="5F6B78")
MONEY = '#,##0.00;[Red]-#,##0.00'
THIN = Side(style="thin", color="D9DEE4")


def _header(ws, row, labels, widths=None):
    for i, lab in enumerate(labels, 1):
        c = ws.cell(row=row, column=i, value=lab)
        c.fill, c.font = HEAD_FILL, HEAD_FONT
        c.alignment = Alignment(vertical="center", wrap_text=True)
    if widths:
        for i, w in enumerate(widths, 1):
            ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = ws.cell(row=row + 1, column=1)


def _money_cols(ws, first_row, cols):
    for r in range(first_row, ws.max_row + 1):
        for c in cols:
            cell = ws.cell(row=r, column=c)
            if isinstance(cell.value, (int, float)):
                cell.number_format = MONEY


def build_workbook(destination: Path, months: dict[str, dict], cur_month: str,
                   changes: list[Change], compared_to: str | None) -> Path:
    cur = months[cur_month]
    wb = Workbook()

    # ---- What changed
    ws = wb.active
    ws.title = "What changed"
    ws["A1"] = f"{cur['month_label']} — what's worth a look"
    ws["A1"].font = TITLE_FONT
    if compared_to:
        ws["A2"] = f"Compared with {compared_to}. Red rows need attention; yellow rows are differences worth knowing about."
    else:
        ws["A2"] = ("First month on record, so there is nothing to compare against yet. "
                    "Run last month's statement through the app and this sheet fills in. "
                    "The checks below don't need a prior month.")
    ws["A2"].font = SUB_FONT
    _header(ws, 4, ["Property", "What", "Detail", "This month", "Last month"], [34, 26, 90, 14, 14])
    r = 5
    if not changes:
        ws.cell(row=r, column=1, value="Nothing stood out this month.")
    for c in changes:
        fill = ATTN_FILL if c.severity == "attention" else NOTE_FILL
        vals = [c.property, c.what, c.detail, c.this, c.last]
        for i, v in enumerate(vals, 1):
            cell = ws.cell(row=r, column=i, value=v)
            cell.fill = fill
            cell.alignment = Alignment(wrap_text=True, vertical="top")
            if isinstance(v, (int, float)) and i >= 4:
                cell.number_format = MONEY
        r += 1

    # ---- This month
    ws = wb.create_sheet("This month")
    ws["A1"] = f"{cur['month_label']} — every property at a glance"
    ws["A1"].font = TITLE_FONT
    cols = ["Property", "Tenant", "Status", "Rent roll (rent + recurring)", "Rent received", "Other income",
            "Total income", "Mgmt fee", "Mgmt %", "Other expenses", "Total expenses", "Owner payment",
            "Net income (income stmt)", "Net income YTD", "Beginning cash", "Ending cash", "Bills due", "Past due"]
    widths = [34, 22, 10, 14, 13, 12, 13, 11, 8, 13, 13, 13, 14, 14, 13, 13, 11, 11]
    _header(ws, 3, cols, widths)
    r = 4
    tot = [0.0] * len(cols)
    for name in sorted(cur["properties"]):
        f = figures(cur["properties"][name])
        row = [name, f["tenant"], f["status"], f["scheduled_rent"], f["rent_received"], f["other_income"],
               f["total_income"], f["mgmt_fee"], f["mgmt_pct"], f["other_expense"], f["total_expense"],
               f["owner_payment"], f["net_income"], f["net_income_ytd"], f["beginning_cash"], f["ending_cash"],
               f["bills_due"], f["past_due"]]
        for i, v in enumerate(row, 1):
            cell = ws.cell(row=r, column=i, value=v)
            if isinstance(v, (int, float)) and i != 9:
                cell.number_format = MONEY
                tot[i - 1] += v
            elif i == 9 and isinstance(v, (int, float)):
                cell.number_format = '0.0"%"'
        r += 1
    ws.cell(row=r, column=1, value="Total").font = Font(bold=True)
    for i in range(4, len(cols) + 1):
        if i == 9:
            continue
        cell = ws.cell(row=r, column=i, value=round(tot[i - 1], 2))
        cell.number_format, cell.font = MONEY, Font(bold=True)

    # ---- History
    ws = wb.create_sheet("History")
    ws["A1"] = "Month by month"
    ws["A1"].font = TITLE_FONT
    ws["A2"] = "One block per figure; properties down the side, months across."
    ws["A2"].font = SUB_FONT
    month_keys = sorted(months)
    labels = [months[m]["month_label"] for m in month_keys]
    all_props = sorted({p for m in months.values() for p in m["properties"]})
    r = 4
    for title, key in [("Rent received", "rent_received"), ("Total expenses", "total_expense"),
                       ("Net income (income statement)", "net_income"), ("Owner payment", "owner_payment"),
                       ("Ending cash", "ending_cash"), ("Past due", "past_due")]:
        ws.cell(row=r, column=1, value=title).font = Font(bold=True, size=12)
        r += 1
        for i, lab in enumerate(["Property"] + labels, 1):
            c = ws.cell(row=r, column=i, value=lab)
            c.fill, c.font = HEAD_FILL, HEAD_FONT
        r += 1
        for name in all_props:
            ws.cell(row=r, column=1, value=name)
            for j, m in enumerate(month_keys, 2):
                p = months[m]["properties"].get(name)
                v = figures(p)[key] if p else None
                cell = ws.cell(row=r, column=j, value=v)
                if isinstance(v, (int, float)):
                    cell.number_format = MONEY
            r += 1
        r += 1
    ws.column_dimensions["A"].width = 34
    for j in range(2, len(labels) + 2):
        ws.column_dimensions[get_column_letter(j)].width = 15
    ws.freeze_panes = "B5"

    # ---- Ledger
    ws = wb.create_sheet("Ledger")
    _header(ws, 1, ["Month", "Property", "Date", "Payee / Payer", "Type", "Reference", "Description",
                    "Income", "Expense", "Balance"], [11, 34, 11, 26, 14, 11, 60, 12, 12, 12])
    r = 2
    for m in reversed(month_keys):
        md = months[m]
        for name in sorted(md["properties"]):
            for l in md["properties"][name].get("ledger") or []:
                vals = [md["month_label"], name, l.get("date"), l.get("payee"), l.get("type"), l.get("reference"),
                        l.get("description"), l.get("income"), l.get("expense"), l.get("balance")]
                for i, v in enumerate(vals, 1):
                    ws.cell(row=r, column=i, value=v)
                r += 1
    _money_cols(ws, 2, [8, 9, 10])
    ws.auto_filter.ref = f"A1:J{max(r - 1, 1)}"

    out = Path(destination) / SUMMARY_DIR / WORKBOOK
    out.parent.mkdir(parents=True, exist_ok=True)
    tmp = out.with_suffix(".xlsx.partial")
    wb.save(tmp)
    tmp.replace(out)
    return out


# ------------------------------------------------------------------ entry point

class NotSummarised(Exception):
    """Raised when a file shouldn't go into the history (with a plain-English reason)."""


def update_summary(destination: Path, plan, reader) -> SummaryResult:
    """Read this month's figures, store them, rebuild the workbook, report changes.

    The comparison shown is always for the LATEST month on record against the
    month before it — so running earlier months to fill in history never hides
    the current month's findings.
    """
    destination = Path(destination)
    if getattr(plan, "month_guessed", False):
        raise NotSummarised("The statement period couldn't be read from this file, so its figures "
                            "were not added to the Monthly Summary (they'd be filed under the wrong month).")
    if not any(s.code for s in plan.sections):
        raise NotSummarised("No property sections were recognised, so nothing was added to the Monthly Summary.")
    month_data = collect_month(plan, reader)
    save_month(destination, month_data)
    months = load_months(destination)
    months[month_data["month"]] = month_data
    latest = max(months)
    earlier = [m for m in sorted(months) if m < latest]
    prev = months[earlier[-1]] if earlier else None
    changes = compare(months[latest], prev)
    wb = build_workbook(destination, months, latest, changes, prev["month_label"] if prev else None)
    return SummaryResult(workbook=wb, month=latest, month_label=months[latest]["month_label"],
                         compared_to=prev["month_label"] if prev else None, changes=changes,
                         months_on_record=len(months), ran_month_label=month_data["month_label"],
                         ran_is_latest=(latest == month_data["month"]))


def describe(result: SummaryResult) -> str:
    lines = []
    if result.compared_to:
        lines.append(f"Compared {result.month_label} with {result.compared_to}: "
                     f"{len(result.attention)} to look at, {len(result.notes)} smaller differences.")
    else:
        lines.append(f"{result.month_label} is the first month on record, so there is no prior month to compare "
                     f"against yet. Run an earlier statement and the comparison fills in.")
    for c in result.changes:
        mark = "!!" if c.severity == "attention" else " -"
        lines.append(f"  {mark} {c.property}: {c.detail}")
    lines.append(f"Workbook: {result.workbook}")
    return "\n".join(lines)


if __name__ == "__main__":
    import sys
    from pypdf import PdfReader
    from splitter import plan_split
    src, dest = Path(sys.argv[1]), Path(sys.argv[2])
    res = update_summary(dest, plan_split(src), PdfReader(str(src)))
    print(describe(res))
