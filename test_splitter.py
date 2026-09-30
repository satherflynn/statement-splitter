"""Regression test: split the fictional months, then check the comparison.

The sample is billed like the real statements (water monthly, garbage
quarterly, property tax in August), so this test fails if quarterly bills
ever start cluttering "What changed" again.

Run:  python test_splitter.py
"""
import tempfile
from pathlib import Path

from openpyxl import load_workbook
from pypdf import PdfReader

import make_sample_packet
from ledger import read_property_month
from splitter import plan_split, write_split
from summary import update_summary

EXPECTED_AUGUST = [  # (folder name, page count)
    ("004-01 123 Main Street", 4),
    ("004-02 125 Main Street", 4),
    ("004-03 127 Main Street", 4),
    ("004-05 4410 Meadow Lane", 4),
    ("004-06 4412 - 4412 Meadow Lane", 4),
    ("004-07 9 Very Long Address Name Court", 4),
    ("004-01 123 Main Street (duplicate)", 3),
    ("Unsorted pages", 2),                    # bill copies appended after the last property
]

# (property, "what") that August MUST flag as needing attention
EXPECTED_ATTENTION = {
    ("004-02 125 Main Street", "Rent roll changed"),
    ("004-02 125 Main Street", "Tenant changed"),
    ("004-03 127 Main Street", "Management fee out of line"),
    ("004-06 4412 - 4412 Meadow Lane", "Tenant past due"),
    ("004-06 4412 - 4412 Meadow Lane", "Rent received changed"),
    ("004-07 9 Very Long Address Name Court", "Negative cash balance"),
    ("004-07 9 Very Long Address Name Court", "Reversals don't net to zero"),
    ("004-01 123 Main Street (duplicate)", "Negative cash balance"),
    ("004-01 123 Main Street (duplicate)", "Unit not current"),
}
# ...and must note as smaller differences
EXPECTED_NOTES = {
    ("004-01 123 Main Street", "Bill amount changed"),        # garbage $74.10 vs $71.04 paid in May
    ("004-05 4410 Meadow Lane", "Expected bill not paid"),     # monthly water missing
    ("004-05 4410 Meadow Lane", "New charge"),                 # plumbing repair
    ("004-01 123 Main Street (duplicate)", "New property"),
}


def _month(month, bills):
    """A minimal month on record for compare(): one property, just its bills."""
    ledger = [{"date": f"{month[5:]}/15/{month[:4]}", "payee": payee, "type": typ, "reference": "",
               "description": desc, "income": None, "expense": amt, "balance": None}
              for payee, typ, desc, amt in bills]
    return {"month": month, "month_label": month, "properties": {"004-01 Test": {"ledger": ledger, "units": []}}}


def check_irregular_bills():
    """Sep 2026 feedback: property tax was reported as 'usually paid every 2
    months'. Tax installments are irregular and the reserve postings are
    bookkeeping; neither may be reported as a missing bill. A monthly bill
    that stops must still be reported."""
    from summary import compare
    tax = ("Washoe County Treasurer", "Check", "Property Tax - 50606103 Jul", 111.09)
    reserve = ("Property Tax Reserve", "Check", "Property Tax - July 2026", 32.29)
    water = ("Somewhere Water Co.", "Payment", "Water - 040160-000", 61.25)
    months = [_month("2026-05", [tax, reserve, water]), _month("2026-06", [reserve, water]),
              _month("2026-07", [tax, reserve, water]), _month("2026-08", [reserve, water])]
    sept = _month("2026-09", [])      # nothing paid in September
    missing = [c for c in compare(sept, months) if c.what == "Expected bill not paid"]
    assert [c.detail.split(":")[0] for c in missing] == ["Somewhere Water Co."], [c.detail for c in missing]
    assert not any("Property Tax Reserve" in c.detail for c in compare(sept, months))


def run_all(paths, dest):
    return [update_summary(dest, plan_split(p), PdfReader(str(p))) for p in paths]


def main():
    check_irregular_bills()
    with tempfile.TemporaryDirectory() as tmp:
        tmp = Path(tmp)
        may, jun, jul, aug = make_sample_packet.build(tmp / "samples")
        dest = tmp / "out"

        # --- split
        plan = plan_split(aug)
        got = [(s.folder_name, s.page_count) for s in plan.sections]
        assert plan.month == "2026-08" and plan.month_label == "August 2026", (plan.month, plan.month_label)
        assert got == EXPECTED_AUGUST, "\n".join(f"  {g}  vs  {e}" for g, e in zip(got, EXPECTED_AUGUST)) + f"\n({len(got)} sections)"
        assert sum(n for _, n in got) == plan.total_pages
        written = write_split(plan, dest)
        for w in written:
            expect = "2026-08.pdf" if w.section.code else f"2026-08 pages {w.section.first_page}-{w.section.last_page}.pdf"
            assert w.path.name == expect, w.path
            assert len(PdfReader(str(w.path)).pages) == w.section.page_count
        assert any("more than once" in w for w in plan.warnings), plan.warnings
        assert any("isn't one of AppFolio's reports" in w for w in plan.warnings), plan.warnings

        # --- summary, months in order
        r_may, r_jun, r_jul, r_aug = run_all([may, jun, jul, aug], dest)
        assert r_may.compared_to is None and r_may.months_on_record == 1
        # Quiet months must be QUIET: garbage paid in May and not June/July is its rhythm, not news.
        for r in (r_may, r_jun, r_jul):
            assert not r.changes, (r.month_label, [c.detail for c in r.changes])
        assert r_aug.compared_to == "July 2026" and r_aug.months_on_record == 4

        attn = {(c.property, c.what) for c in r_aug.attention}
        notes = {(c.property, c.what) for c in r_aug.notes}
        assert EXPECTED_ATTENTION <= attn, f"not raised: {EXPECTED_ATTENTION - attn}\nraised: {sorted(attn)}"
        assert EXPECTED_NOTES <= notes, f"not raised: {EXPECTED_NOTES - notes}\nraised: {sorted(notes)}"
        # The first property-tax run and the quarterly garbage run are routine, not news.
        noisy = [c.detail for c in r_aug.changes if "Washoe" in c.detail
                 or ("Waste Management" in c.detail and c.property != "004-01 123 Main Street")]
        assert not noisy, noisy
        garbage = next(c for c in r_aug.notes if c.what == "Bill amount changed")
        assert "last paid $71.04 in May 2026" in garbage.detail, garbage.detail
        assert not any(c.what == "Owner payment changed a lot" for c in r_aug.changes)
        assert len(r_aug.notes) <= 8, [c.detail for c in r_aug.notes]
        assert not [c for c in r_aug.attention if c.property == "004-01 123 Main Street"]

        # --- the This month sheet is two stacked tables with Property + Tenant on each
        ws = load_workbook(r_aug.workbook)["This month"]
        rows = [[c.value for c in row] for row in ws.iter_rows(min_row=1, max_row=ws.max_row)]
        headers = [r for r in rows if r[0] == "Property"]
        assert len(headers) == 2, headers
        assert headers[0][1] == headers[1][1] == "Tenant"
        assert headers[0][-1] == "Total expenses" or "Total expenses" in headers[0], headers[0]
        assert headers[1][2] == "Owner payment", headers[1]
        assert ws.max_column <= 11, ws.max_column

        # --- out of order: August first, then July — the comparison stays on August
        dest2 = tmp / "out2"
        r1 = update_summary(dest2, plan_split(aug), PdfReader(str(aug)))
        assert r1.compared_to is None and r1.ran_is_latest
        r2 = update_summary(dest2, plan_split(jul), PdfReader(str(jul)))
        assert r2.ran_month_label == "July 2026" and not r2.ran_is_latest
        assert r2.month_label == "August 2026" and r2.compared_to == "July 2026"
        assert {(c.property, c.what) for c in r2.attention} >= EXPECTED_ATTENTION

        # --- every property's ledger reconciles to the cent, every month
        for path in (may, jun, jul, aug):
            reader, p = PdfReader(str(path)), plan_split(path)
            for s in p.sections:
                if s.code is None:
                    continue
                d = read_property_month(reader, s.first_page, s.last_page)
                inc = sum(l["income"] or 0 for l in d["ledger"]); exp = sum(l["expense"] or 0 for l in d["ledger"])
                assert abs(d["beginning_cash"] + inc - exp - d["ending_cash"]) < 0.011, (path.name, s.folder_name)

    print(f"OK — August split into {len(EXPECTED_AUGUST)} sections / {plan.total_pages} pages; "
          f"May–July quiet; August raised {len(r_aug.attention)} to look at + {len(r_aug.notes)} smaller differences; "
          f"ledgers reconcile")


if __name__ == "__main__":
    main()
